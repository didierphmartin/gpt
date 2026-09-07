"""Port of backend/src/AgentTeam/Services/PythonEmitHelpers.php (756 lines).

Shared Python-code-emission helpers used by LangGraphGenerator and (from
Task 2/3 of this port) ADKGenerator. All methods are pure static; no DB or
state dependencies.

Methods moved verbatim from LangGraphGenerator (PHP docblock, lines 15-19):
  - jsonToPython()            PHP value -> Python literal (True/False/None)
  - mcpClientBlock()          MCP JSON-RPC client Python functions
  - skillDepsBlock()          SKILLS_DIR globals + _ensure_skill_deps Python function
  - documentConverterBlock()  _convert_doc_to_markdown() Python function

New helpers (not previously in LangGraphGenerator):
  - pyStr()          PHP string -> Python double-quoted string literal

The four `*Block()` methods return LITERAL generated-Python-code text (the
scripts LangGraph/ADK/MAF/NOOA generators emit). That text must be
byte-identical to the PHP nowdoc content -- the constants above are
extracted programmatically from PythonEmitHelpers.php (no manual
transcription) and pinned by test_python_emit_helpers_pin.py.
"""
from __future__ import annotations

import json
import re
import unicodedata

from app.support.phpcompat import php_date, php_empty, php_intval, php_strval, php_trim
from app.support.phpjson import dumps_pretty

# PHP `preg_match('/\((\S+)\)$/', $p, $m)` (no `/u` modifier -> ASCII \S, not
# PCRE_UCP -- re.ASCII mirrors that).
_PARENT_ID_RE = re.compile(r'\((\S+)\)$', re.ASCII)


def _coalesce(*vals):
    """PHP `??` (null-coalescing) chain: first argument that is not None,
    else None. NEVER use bare `or` for a `??` port -- `or` also falls
    through on '', 0, False, [], {} which `??` does not."""
    for v in vals:
        if v is not None:
            return v
    return None


def _phpFloatToken(v: float) -> str:
    """PHP `json_encode((float) $v)` under `serialize_precision=-1` (the
    effective runtime precision at these call sites -- see
    `_jsonEncodeUnescapedUnicode` docstring below): shortest round-trip
    digits, but a WHOLE-NUMBER float prints WITHOUT a trailing '.0' --
    `json_encode((float) 1.0)` is `"1"`, not Python `json.dumps`'s `"1.0"`.
    Verified against `php -r` 2026-09-07: 0.0/1.0/2.0/100.0/-1.0/-0.0 ->
    '0'/'1'/'2'/'100'/'-1'/'-0'; 0.7/2.5/0.1/1.23456789012345 unchanged.
    """
    s = repr(v)
    if s.endswith('.0'):
        s = s[:-2]
    return s


def _jsonEncodeUnescapedUnicode(value) -> str:
    """`json_encode($v, JSON_UNESCAPED_UNICODE)` (PHP 587, 596, 601, 735, 742):
    literal (non-escaped) unicode, but slashes ARE escaped (JSON_UNESCAPED_SLASHES
    is NOT set at these call sites), compact separators (no JSON_PRETTY_PRINT).

    Float formatting: these calls run after the caller
    (LangGraphGenerator::generate/analyzeForEmit, PHP ~line 692-695) has done
    `ini_set('serialize_precision', '-1')` for shortest round-tripping floats
    ("temperatures like 0.5999999999999999... strict models reject") --
    NOT this environment's php.ini `serialize_precision=100` default, which
    would bake long decimal expansions no caller of these helpers ever
    actually emits. A `float` argument is routed through `_phpFloatToken`
    (Python's `json.dumps` keeps a trailing '.0' PHP's encoder does not);
    every other value type uses plain `json.dumps`, whose string/bool/None
    formatting already matches PHP's under these flags.
    """
    if isinstance(value, float):
        return _phpFloatToken(value)
    return json.dumps(value, ensure_ascii=False, separators=(',', ':')).replace('/', '\\/')


def _mbStrimwidth(s: str, width: int, trim_marker: str = '…') -> str:
    """Port of PHP `mb_strimwidth($s, 0, $width, $trim_marker)` for start=0.

    PHP's mbstring measures DISPLAY width, not character count: East-Asian
    Wide/Fullwidth characters count as 2 columns, everything else (narrow,
    neutral, ambiguous, halfwidth) counts as 1 -- verified against
    `php -r 'mb_strimwidth(...)'` 2026-09-07 for plain ASCII (71 'a' chars ->
    69 'a' + the marker, i.e. the marker's own width is reserved from the
    budget) and repeated wide characters (36 full-width chars -> 34 chars +
    marker). If the string's total width already fits, it is returned
    unchanged with NO marker appended.
    """
    def w(ch: str) -> int:
        return 2 if unicodedata.east_asian_width(ch) in ('W', 'F') else 1

    total = sum(w(c) for c in s)
    if total <= width:
        return s
    marker_w = sum(w(c) for c in trim_marker)
    budget = width - marker_w
    out = []
    acc = 0
    for ch in s:
        cw = w(ch)
        if acc + cw > budget:
            break
        out.append(ch)
        acc += cw
    return ''.join(out) + trim_marker



# --- literal generated-Python-code blocks, extracted byte-exact from ---
# --- backend/src/AgentTeam/Services/PythonEmitHelpers.php nowdocs    ---

_MCP_CLIENT_BLOCK = r'''def _normalize_mcp_url(url: str) -> str:
    """Ensure the URL points to the MCP endpoint.

    - If it already ends with /mcp, leave it alone.
    - If it ends with a .php file, it IS the endpoint already (XAMPP-style
      MCP server scripts like mcp-server.php). Don't append anything.
    - Otherwise, append /mcp per the MCP convention.
    """
    url = url.rstrip("/")
    if url.endswith("/mcp") or url.endswith(".php"):
        return url
    return url + "/mcp"

_MCP_ENDPOINTS: dict = {}   # registered server URL -> endpoint that accepted the handshake

def _init_mcp_session(url: str, headers: dict) -> str | None:
    """Initialize an MCP session with the server.

    The MCP protocol requires a handshake before tool calls:
    1. Client sends 'initialize' with protocol version and capabilities
    2. Server responds with its capabilities
    3. Client sends 'notifications/initialized' to confirm

    Tries the /mcp-normalised URL first, then the URL exactly as registered
    (servers mounted at their bare URL, e.g. PHP mock servers, answer there
    and 404 or error on /mcp). The endpoint that answers is remembered.
    Returns that endpoint URL, or None when no candidate accepted the handshake.
    """
    init_req = {
        "jsonrpc": "2.0", "id": 1, "method": "initialize",
        "params": {
            "protocolVersion": "2024-11-05",
            "clientInfo": {"name": "LangGraph-Workflow", "version": "1.0.0"},
            "capabilities": {}
        }
    }
    candidates = [_MCP_ENDPOINTS[url]] if url in _MCP_ENDPOINTS else []
    for cand in (_normalize_mcp_url(url), url):
        if cand not in candidates:
            candidates.append(cand)
    last_err = None
    for mcp_url in candidates:
        try:
            with httpx.Client(timeout=30) as c:
                r = c.post(mcp_url, json=init_req, headers=headers)
                r.raise_for_status()
                data = _parse_mcp_response(r.text)
                if data is None or "error" in (data or {}):
                    last_err = (data or {}).get("error") if data else "invalid response"
                    continue
                # Send initialized notification
                c.post(mcp_url, json={
                    "jsonrpc": "2.0",
                    "method": "notifications/initialized",
                    "params": {}
                }, headers=headers)
                _MCP_ENDPOINTS[url] = mcp_url
                return mcp_url
        except Exception as e:
            last_err = e
    print(f"[warn] MCP init failed for {url}: {last_err}")
    return None

def _parse_mcp_response(text: str) -> dict | None:
    """Parse a JSON-RPC response from an MCP server.

    MCP servers may respond in two formats:
    - Plain JSON: standard JSON-RPC response body
    - SSE (Server-Sent Events): lines prefixed with 'data:' containing JSON
    This function tries plain JSON first, then falls back to SSE parsing.
    """
    import json as _json
    try:
        return _json.loads(text)
    except _json.JSONDecodeError:
        pass
    for line in text.split("\n"):
        line = line.strip()
        if line.startswith("data:"):
            data = line[5:].strip()
            if data:
                try:
                    return _json.loads(data)
                except _json.JSONDecodeError:
                    continue
    return None

def _call_mcp_tool(server_url: str, tool_name: str,
                   arguments: dict) -> str:
    """Call a tool on an MCP server via JSON-RPC 2.0.

    Sequence: init session -> send tools/call -> parse response -> extract text.
    The MCP response contains a 'content' array; we extract all text items
    and join them. If no text is found, falls back to raw JSON.
    """
    headers = {"Content-Type": "application/json",
               "Accept": "application/json, text/event-stream, */*"}

    mcp_url = _init_mcp_session(server_url, headers) or _normalize_mcp_url(server_url)

    request = {
        "jsonrpc": "2.0", "id": int(time.time()),
        "method": "tools/call",
        "params": {"name": tool_name,
                   "arguments": arguments or {}}
    }
    try:
        with httpx.Client(timeout=180) as c:
            r = c.post(mcp_url, json=request, headers=headers)
            r.raise_for_status()
            data = _parse_mcp_response(r.text)
    except Exception as e:
        return json.dumps({"error": str(e)})

    if data is None:
        return json.dumps({"error": "invalid response"})
    if "error" in data:
        return json.dumps({"error": data["error"].get("message", "unknown")})

    content = (data.get("result") or {}).get("content", [])
    texts = [c.get("text", "") for c in content
             if isinstance(c, dict) and c.get("type") == "text"]
    return "\n".join(t for t in texts if t) or json.dumps(data.get("result"))[:8000]
'''


_SKILL_DEPS_BLOCK = r'''# ---------------------------------------------------------------------------
# Local run_skill_script — execute a folder-backed skill's Python script as a
# subprocess. Skills are real CPython (the browser runs them in Pyodide; here
# we shell out). Registered into the catalog at startup so skill-bound agents
# actually run their skill instead of fabricating output.
# ---------------------------------------------------------------------------
SKILLS_DIR = os.environ.get(
    "SYNERGYAI_SKILLS_DIR",
    os.path.join(os.path.expanduser("~"), "Documents", "synergyAI", "skills"),
)

# Skills declare their PyPI deps in SKILL.md frontmatter
# (e.g. `dependencies: [beautifulsoup4]`). The browser installs them via
# micropip; here we pip-install missing ones into the runner venv on first
# use. Cached + lock-guarded because parallel branches call skills at once.
_INSTALLED_DEPS = set()
_DEPS_LOCK = threading.Lock()


def _ensure_skill_deps(skill_dir: str) -> None:
    md = os.path.join(skill_dir, "SKILL.md")
    if not os.path.isfile(md):
        return
    deps = []
    try:
        lines = open(md, encoding="utf-8").read().splitlines()
    except Exception:
        return
    # Only the YAML frontmatter (between leading '---' fences) is authoritative;
    # a 'dependencies:' line in the prose body is not a real declaration.
    if not lines or lines[0].strip() != "---":
        return
    for line in lines[1:]:
        s = line.strip()
        if s == "---":
            break
        if s.lower().startswith("dependencies:"):
            val = s.split(":", 1)[1].strip().strip("[]")
            deps = [d.strip().strip("'").strip('"') for d in val.split(",") if d.strip()]
            break
    with _DEPS_LOCK:
        for d in deps:
            if d in _INSTALLED_DEPS:
                continue
            _INSTALLED_DEPS.add(d)
            try:
                print(f"  [run_skill_script] ensuring dependency: {d}", flush=True)
                r = subprocess.run([sys.executable, "-m", "pip", "install", "-q", d],
                                   capture_output=True, text=True, timeout=600)
                if r.returncode != 0:
                    print(f"  [run_skill_script] pip install {d} FAILED: "
                          f"{(r.stderr or '')[-300:]}", flush=True)
            except Exception as e:
                print(f"  [run_skill_script] pip install {d} error: {e}", flush=True)
'''


_SKILL_FS_SYNC_BLOCK = r'''

# Skill filesystem (mirrors the browser interpreter + ADK): real bucketed /outputs and
# /scratch dirs, virtual-path remap, env exposure, and a deliverable stash so a skill's
# produced file (not the model's chatter) can become the node output.
SKILL_OUTPUTS_ROOT = os.environ.get("SYNERGYAI_OUTPUT_ROOT") or os.path.join(os.path.dirname(SKILLS_DIR), "outputs")
SKILL_SCRATCH_DIR = os.environ.get("SYNERGYAI_SCRATCH_DIR") or os.path.join(os.path.dirname(SKILLS_DIR), "scratch")
_LAST_SKILL_OUTPUTS = {}


def _skill_output_dir(dir_name: str) -> str:
    group = dir_name.split("/")[0] if "/" in dir_name else ""
    if group and all(c.isalnum() or c in "._-" for c in group):
        return os.path.join(SKILL_OUTPUTS_ROOT, group)
    return SKILL_OUTPUTS_ROOT


def _remap_virtual_path(p, out_dir: str) -> str:
    p = str(p)
    for virt, real in (("/outputs", out_dir), ("/scratch", SKILL_SCRATCH_DIR)):
        if p == virt:
            return real
        if p.startswith(virt + "/"):
            return os.path.join(real, p[len(virt) + 1:])
    return p


def _fix_overescaped(text):
    """Repair JSON-style over-escaping in model-carried document content.

    Models (Claude especially) sometimes emit large documents inside tool
    arguments as an escaped single-line string: literal \\n / \\" sequences
    and ZERO real newlines. Written verbatim, the HTML/CSS is broken and the
    page renders unstyled (a live report shipped with 1131 literal \\n).
    Mirrors the platform chat's _coerceInputFiles over-escape recovery.
    Detection is conservative: many escape sequences AND no real newlines."""
    s = str(text)
    if s.count("\\n") > 5 and s.count("\n") == 0:
        print("  [stage] repairing over-escaped content ("
              + str(s.count("\\n")) + " literal escape sequences)", flush=True)
        s = (s.replace("\\\\", "\x00").replace("\\n", "\n").replace("\\t", "\t")
              .replace("\\r", "").replace('\\"', '"').replace("\\'", "'").replace("\x00", "\\"))
    return s


def _run_skill_script(dir_name: str, script: str, argv=None,
                      input_files=None, read_outputs=None) -> str:
    if isinstance(argv, str):
        try:
            argv = json.loads(argv)
        except Exception:
            argv = [argv]
    argv = list(argv) if argv else []
    skill_dir = os.path.join(SKILLS_DIR, *str(dir_name).split("/"))
    script_path = os.path.join(skill_dir, *str(script).split("/"))
    if not os.path.isfile(script_path):
        return f"ERROR: skill script not found: {dir_name}/{script} (looked in {script_path})"
    _ensure_skill_deps(skill_dir)
    out_dir = _skill_output_dir(dir_name)
    os.makedirs(out_dir, exist_ok=True)
    os.makedirs(SKILL_SCRATCH_DIR, exist_ok=True)
    # Remap the interpreter's virtual /outputs & /scratch in argv to real dirs.
    argv = [_remap_virtual_path(a, out_dir) for a in argv]
    # Stage input_files (e.g. HTML the skill will render) at their remapped paths.
    if isinstance(input_files, dict):
        for raw_path, content in input_files.items():
            try:
                target = _remap_virtual_path(raw_path, out_dir)
                os.makedirs(os.path.dirname(target) or ".", exist_ok=True)
                with open(target, "w", encoding="utf-8") as fh:
                    fh.write(_fix_overescaped(content))
            except Exception as e:
                print(f"  [run_skill_script] could not stage {raw_path}: {e}")
    env = dict(
        os.environ,
        SYNERGYAI_OUTPUT_DIR=out_dir,
        SYNERGYAI_SCRATCH_DIR=SKILL_SCRATCH_DIR,
        SYNERGYAI_SKILL_DIR_NAME=dir_name,
        SYNERGYAI_SKILL_GROUP=(dir_name.split("/")[0] if "/" in dir_name else ""),
    )
    cmd = [sys.executable, script_path] + [str(a) for a in argv]
    print(f"  [run_skill_script] → {dir_name}/{script} argv={argv}", flush=True)
    _t0 = time.monotonic()
    try:
        proc = subprocess.run(cmd, cwd=skill_dir, env=env, capture_output=True,
                              text=True, timeout=300)
    except subprocess.TimeoutExpired:
        print(f"  [run_skill_script] ← {dir_name}/{script}: TIMEOUT after 300s", flush=True)
        return f"ERROR: skill {dir_name}/{script} timed out after 300s"
    _dt = time.monotonic() - _t0
    out = proc.stdout or ""
    _stderr = (proc.stderr or "").strip()
    _head = out.replace(chr(10), " ")[:160]
    print(f"  [run_skill_script] ← {dir_name}/{script}: exit={proc.returncode}, "
          f"{len(out)} chars stdout, {_dt:.1f}s | {_head}", flush=True)
    if proc.returncode != 0 and _stderr:
        print(f"  [run_skill_script]   stderr tail: {_stderr[-400:]}", flush=True)
    if proc.returncode != 0:
        out = (out + f"\n[run_skill_script exit {proc.returncode}]\n"
               + _stderr[-2000:]).strip()
    # Read back requested outputs AND stash them so the skill step can use the produced
    # document as the node output.
    _produced = []
    if isinstance(read_outputs, list):
        for rel in read_outputs:
            try:
                with open(_remap_virtual_path(rel, out_dir), "r", encoding="utf-8") as fh:
                    _content = fh.read()
                out += f"\n\n[output file {rel}]\n" + _content
                _produced.append(_content)
            except Exception:
                pass
    if _produced:
        _LAST_SKILL_OUTPUTS[dir_name] = _produced
    return out or "(skill produced no stdout)"


def _read_skill_md(dir_name: str) -> str:
    """The live SKILL.md body (progressive disclosure); frontmatter stripped."""
    path = os.path.join(SKILLS_DIR, *str(dir_name).split("/"), "SKILL.md")
    try:
        with open(path, "r", encoding="utf-8") as f:
            text = f.read()
    except OSError:
        return f"(SKILL.md not found for skill '{dir_name}')"
    if text.startswith("---"):
        end = text.find("\n---\n", 3)
        if end != -1:
            nl = text.find("\n", end + 1)
            text = text[nl + 1:] if nl != -1 else ""
    return text.strip()
'''


_DOCUMENT_CONVERTER_BLOCK = r'''def _convert_doc_to_markdown(path: str) -> str:
    """Read a file and return its contents as Markdown.

    Text-native formats (HTML/MD/TXT/JSON/YAML/CSV/XML/source) are returned
    verbatim. Binary office formats and PDFs are routed to their matching
    pure-Python parser. Imports are lazy so the script doesn't pull in a
    library unless the corresponding format is actually attached.

    Raises FileNotFoundError if the path doesn't exist, ValueError for an
    unsupported extension, ImportError if the format's library is missing.
    """
    from pathlib import Path
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"document not found: {path}")
    ext = p.suffix.lower()

    text_native = {
        ".html", ".htm", ".md", ".markdown", ".txt", ".text",
        ".json", ".yaml", ".yml", ".xml", ".csv", ".tsv",
        ".css", ".scss", ".less",
        ".js", ".mjs", ".ts", ".tsx", ".jsx",
        ".py", ".rb", ".go", ".rs", ".java", ".c", ".cc", ".cpp", ".h", ".hpp",
        ".sh", ".bash", ".zsh", ".sql", ".log", ".rtf", ".ini", ".toml",
    }
    if ext in text_native:
        return p.read_text(encoding="utf-8", errors="replace")

    if ext == ".docx":
        import mammoth
        with p.open("rb") as f:
            return mammoth.convert_to_markdown(f).value

    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(str(p))
        sections = [f"# {p.name}", f"_{len(prs.slides)} slide(s)_"]
        for i, slide in enumerate(prs.slides, start=1):
            layout = slide.slide_layout.name if slide.slide_layout else "?"
            sections.append(f"## Slide {i} -- {layout}")
            title_text = ""
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                title_text = slide.shapes.title.text_frame.text.strip()
            if title_text:
                sections.append(f"### {title_text}")
            body = []
            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue
                if not getattr(shape, "has_text_frame", False):
                    continue
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text or "" for run in para.runs).strip()
                    if not line and para.text:
                        line = para.text.strip()
                    if line:
                        indent = "  " * (para.level or 0)
                        body.append(f"{indent}- {line}")
            if body:
                sections.append("\n".join(body))
            if slide.has_notes_slide:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()
                if notes:
                    sections.append("### Notes")
                    sections.append(notes)
        return "\n\n".join(sections) + "\n"

    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(str(p), data_only=True)
        sections = [f"# {p.name}", f"_{len(wb.sheetnames)} sheet(s)_"]
        max_rows = 200
        for name in wb.sheetnames:
            ws = wb[name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                row_vals = list(row)
                while row_vals and (row_vals[-1] is None or row_vals[-1] == ""):
                    row_vals.pop()
                if row_vals or not rows:
                    rows.append(row_vals)
            while rows and not any(c not in (None, "") for c in rows[-1]):
                rows.pop()
            sections.append(f"## Sheet: {name}")
            sections.append(f"_{len(rows)} row(s) -- {ws.max_column} col(s)_")
            if not rows:
                sections.append("_(empty)_")
                continue
            width = max(len(r) for r in rows)
            rows = [list(r) + [""] * (width - len(r)) for r in rows]
            truncated = len(rows) > max_rows
            if truncated:
                rows = rows[:max_rows]
            def cell(v):
                return ("" if v is None else str(v)).replace("|", "\\|").replace("\n", " ")
            table = ["| " + " | ".join(cell(v) for v in rows[0]) + " |",
                     "| " + " | ".join(["---"] * width) + " |"]
            for r in rows[1:]:
                table.append("| " + " | ".join(cell(v) for v in r) + " |")
            if truncated:
                table.append(f"\n_(showing first {max_rows} rows)_")
            sections.append("\n".join(table))
        return "\n\n".join(sections) + "\n"

    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(p))
        parts = [f"# {p.name} -- {len(reader.pages)} page(s)"]
        for i, page in enumerate(reader.pages, start=1):
            try:
                t = (page.extract_text() or "").strip()
            except Exception as e:
                t = ""
            parts.append(f"## Page {i}")
            parts.append(t if t else "_(no extractable text)_")
        return "\n\n".join(parts) + "\n"

    raise ValueError(f"unsupported document format: {ext}")
'''


class PythonEmitHelpers:
    """Static-only class -- every method is `@staticmethod`, matching the PHP
    class (PHP 24-756: `class PythonEmitHelpers { public static function ... }`)."""

    @staticmethod
    def jsonToPython(value, forceObject: bool = False) -> str:
        """PHP 35-47.

        json_encode() produces 'true'/'false'/'null', not valid Python. This
        converts to 'True'/'False'/'None'. Uses dumps_pretty() (4-space
        indent, matching Python's own json.dumps(indent=4) output and PHP's
        JSON_PRETTY_PRINT byte-for-byte -- see phpjson.dumps_pretty docstring)
        with unescaped unicode AND unescaped slashes (JSON_UNESCAPED_UNICODE
        | JSON_UNESCAPED_SLASHES), matching the PHP call's flags exactly.

        `forceObject`: PHP `(object) $value` casts an array to stdClass so an
        EMPTY PHP array encodes as `{}` instead of `[]`, and a list-shaped
        array encodes as an object with string-numeric keys instead of a JSON
        array. Python already distinguishes dict (object) from list (array)
        at the value level, so the only case this needs to act on is a caller
        passing a `list` (Python's PHP-array-shaped default for an empty/
        list-shaped array) with forceObject=True: re-key it "0","1","2",...
        exactly like PHP's (object) cast would, so it renders as an object.

        NOTE (ported verbatim, not fixed): PHP replaces the LITERAL substrings
        ': true' / ': false' / ': null' -- so a bare boolean/null that is a
        plain array ELEMENT (no preceding "key: ") is NOT rewritten and would
        stay invalid Python. This is the PHP source's own behaviour; every
        current call site passes an object/dict-shaped value.
        """
        if forceObject and isinstance(value, list):
            value = {str(i): v for i, v in enumerate(value)}
        raw = dumps_pretty(value, unescaped=True)
        raw = raw.replace(': true', ': True')
        raw = raw.replace(': false', ': False')
        raw = raw.replace(': null', ': None')
        return raw

    @staticmethod
    def pyStr(s: str) -> str:
        """PHP 57-65.

        Render a PHP string as a Python double-quoted string literal. Escapes
        backslashes, double-quotes, newlines, carriage returns, and tabs so
        the result is a valid Python string expression.
        """
        s = s.replace('\\', '\\\\')
        s = s.replace('"', '\\"')
        s = s.replace('\n', '\\n')
        s = s.replace('\r', '\\r')
        s = s.replace('\t', '\\t')
        return '"' + s + '"'

    @staticmethod
    def mcpClientBlock() -> str:
        """PHP 75-203 (moved verbatim from LangGraphGenerator::mcpClientBlock()).

        Emit the MCP JSON-RPC client Python block: _normalize_mcp_url,
        _init_mcp_session, _parse_mcp_response, _call_mcp_tool.
        """
        return _MCP_CLIENT_BLOCK

    @staticmethod
    def skillDepsBlock() -> str:
        """PHP 216-275.

        Emit the skill-dependency-install Python block: SKILLS_DIR constant,
        _INSTALLED_DEPS/_DEPS_LOCK globals, _ensure_skill_deps().
        """
        return _SKILL_DEPS_BLOCK

    @staticmethod
    def skillFsSyncBlock() -> str:
        """PHP 292-428.

        Emit the synchronous skill filesystem Python block: SKILL_OUTPUTS_ROOT
        / SKILL_SCRATCH_DIR / _LAST_SKILL_OUTPUTS globals, _skill_output_dir(),
        _remap_virtual_path(), _fix_overescaped(), _run_skill_script(),
        _read_skill_md().
        """
        return _SKILL_FS_SYNC_BLOCK

    @staticmethod
    def documentConverterBlock() -> str:
        """PHP 441-564.

        Emit the document-to-markdown converter Python function
        (_convert_doc_to_markdown).
        """
        return _DOCUMENT_CONVERTER_BLOCK

    # ------------------------------------------------------------------
    # Generated-script documentation (shared by every compile target)
    # ------------------------------------------------------------------

    @staticmethod
    def providerKeyEnv(provider: str) -> str | None:
        """PHP 571-579. Env var that carries the API key of a provider (as the
        runner's .env names them)."""
        return {
            'claude': 'ANTHROPIC_API_KEY', 'anthropic': 'ANTHROPIC_API_KEY',
            'openai': 'OPENAI_API_KEY', 'gemini': 'GOOGLE_API_KEY', 'google': 'GOOGLE_API_KEY',
            'grok': 'XAI_API_KEY', 'xai': 'XAI_API_KEY', 'kimi': 'KIMI_API_KEY', 'moonshot': 'KIMI_API_KEY',
            'deepseek': 'DEEPSEEK_API_KEY', 'glm': 'GLM_API_KEY',
        }.get(provider.lower())

    @staticmethod
    def _nodeSummary(n: dict) -> str:
        """PHP 582-618 (`private static function nodeSummary` -> `_nodeSummary`).
        One-line summary of a node descriptor (see WorkflowGraphAnalyzer.docNodes)."""
        t = n['type']
        if t == 'start':
            p = php_trim(_coalesce(n.get('prompt'), ''))
            if p != '':
                return 'receives the user prompt -- default: ' + _jsonEncodeUnescapedUnicode(_mbStrimwidth(p, 70, '…'))
            return 'receives the user prompt'
        if t == 'output':
            return "merges its parents' outputs verbatim (no LLM) -- the workflow result"
        bits = []
        provider = _coalesce(n.get('provider'), '')
        model = _coalesce(n.get('model'), '')
        if provider != '' or model != '':
            bits.append((provider if not php_empty(provider) else '?') + '/' + (model if model != '' else '(platform default)'))
        if _coalesce(n.get('temperature')) is not None:
            bits.append('temp ' + _jsonEncodeUnescapedUnicode(float(n['temperature'])))
        if _coalesce(n.get('max_tokens')) is not None:
            bits.append('max_tokens ' + str(php_intval(n['max_tokens'])))
        if _coalesce(n.get('thinking')) in ('on', 'off'):
            bits.append('thinking ' + n['thinking'])
        playbook = n['playbook']
        if playbook is not None:
            pb = playbook
            s = 'playbook ' + _jsonEncodeUnescapedUnicode(php_strval(_coalesce(pb.get('title'), '')))
            if _coalesce(pb.get('bound')) is not None:
                s += ", " + php_strval(pb['bound']) + " bound / " + str(php_intval(_coalesce(pb.get('unbound'), 0))) + ' unbound action(s)'
            s += ', writes ' + ('ON' if not php_empty(pb.get('writes')) else 'OFF')
            bits.append(s)
        else:
            bits.append(str(len(n['tools'])) + ' tool(s)')
            if n['skills']:
                bits.append(str(len(n['skills'])) + ' skill(s): ' + ', '.join(n['skills']))
        if n['dispatch']:
            bits.append('DISPATCHER -> one of: ' + ' | '.join(n['dispatch']))
        s = ', '.join(bits)
        if n['dispatch'] and php_empty(n.get('dispatch_supported')):
            s += ' (menu NOT honoured by this target: all children run)'
        if php_empty(n.get('supported')):
            s += " -- NOT RUN BY THIS TARGET (" + t + " nodes are not supported here; the node is skipped)"
        return s

    @staticmethod
    def workflowDocBlock(doc: dict) -> str:
        """PHP 633-722.

        The uniform documentation body every generated script carries in its
        module docstring: provenance, graph nodes, graph edges, execution
        order, target data flow, how to run. Returned WITHOUT the surrounding
        quotes.
        """
        nodes = doc['nodes']
        byId = {}
        for n in nodes:
            byId[n['id']] = n

        def lbl(id_: str) -> str:
            name = byId[id_]['name'] if id_ in byId else 'node'
            return name + ' (' + id_ + ')'

        dispatchers = {}
        for n in nodes:
            if n['dispatch']:
                dispatchers[n['id']] = True
        dispatchOk = not php_empty(doc.get('dispatch_supported'))
        menuNote = ('   (dispatcher menu: one of)' if dispatchOk
                    else '   (dispatcher menu: NOT honoured by this target -- runs in parallel)')
        L: list[str] = []
        L.append('PROVENANCE')
        L.append('==========')
        L.append('  Workflow:   ' + doc['workflow']['name'] + ' (id ' + str(doc['workflow']['id']) + ')')
        L.append('  Target:     ' + doc['target'])
        L.append('  Generated:  ' + php_strval(_coalesce(doc.get('generated_at'), php_date('Y-m-d H:i:s T'))) + ' by the SynergyAI workflow editor')
        L.append('  This file is a frozen snapshot of the workflow. Edits made here are')
        L.append('  overwritten by the next Generate: change the workflow in the editor and')
        L.append('  re-generate instead.')
        L.append('')
        L.append('GRAPH NODES  (id (type): name -- settings; topological order)')
        L.append('===========')
        w = 0
        for n in nodes:
            w = max(w, len(str(n['id']) + ' (' + n['type'] + '):'))
        for n in nodes:
            marker = _coalesce(n.get('marker'), '')
            entry = dict(n)
            entry['dispatch_supported'] = dispatchOk
            head = (str(n['id']) + ' (' + n['type'] + '):').ljust(w)
            L.append('  ' + head + ' ' + n['name'] + ' -- ' + PythonEmitHelpers._nodeSummary(entry)
                      + ('   ' + marker if marker != '' else ''))
        L.append('')
        L.append('GRAPH EDGES  (from -> to)')
        L.append('===========')
        if php_empty(doc.get('edges')):
            L.append('  (none)')
        for e in doc.get('edges') or []:
            f, t = php_strval(e[0]), php_strval(e[1])
            note = menuNote if (f in dispatchers and t in byId and byId[t]['type'] != 'output') else ''
            L.append('  ' + lbl(f) + ' -> ' + lbl(t) + note)
        L.append('')
        L.append('EXECUTION ORDER  (topological layers; nodes on one line run in PARALLEL)')
        L.append('===============')
        for i, layer in enumerate(_coalesce(doc.get('layers'), [])):
            ids = [php_strval(x) for x in layer]
            names = [byId[id_]['name'] if id_ in byId else ('node_' + id_) for id_ in ids]
            line = '  layer ' + str(i) + ':  ' + '  ||  '.join(names)
            viaDispatcher = False
            for id_ in ids:
                for n in nodes:
                    if n['id'] == id_:
                        for p in n['parents']:
                            m = _PARENT_ID_RE.search(p)
                            if m and m.group(1) in dispatchers:
                                viaDispatcher = True
            if len(ids) > 1:
                if viaDispatcher and dispatchOk:
                    line += '   (dispatcher: only the chosen one runs)'
                elif viaDispatcher:
                    line += '   (run in PARALLEL -- dispatcher menu not honoured by this target)'
                else:
                    line += '   (run in PARALLEL)'
            L.append(line)
        L.append('')
        L.append('DATA FLOW')
        L.append('=========')
        for dl in php_strval(_coalesce(doc.get('data_flow'), '')).rstrip('\n').split('\n'):
            L.append(dl)
        L.append('')
        L.append('TO RUN')
        L.append('======')
        run = doc.get('run') if isinstance(doc.get('run'), dict) else {}
        for dep in _coalesce(run.get('deps'), []):
            L.append('  ' + dep)
        keys: dict[str, list[str]] = {}
        for n in nodes:
            if _coalesce(n.get('provider'), '') == '' or php_empty(n.get('supported')):
                continue
            k = PythonEmitHelpers.providerKeyEnv(n['provider'])
            if k is not None:
                keys.setdefault(k, []).append(n['provider'].lower())
        envPath = _coalesce(run.get('env_path'), '')
        envPath = envPath if envPath != '' else '../.env'
        if keys:
            L.append('  # API keys read from ' + envPath + ' for the providers this workflow uses:')
            for k, provs in keys.items():
                seen: list[str] = []
                for p in provs:
                    if p not in seen:
                        seen.append(p)
                L.append('  #   ' + k + ' (' + ', '.join(seen) + ')')
        else:
            L.append('  # No LLM provider is used by a runnable node of this workflow.')
        L.append('  ' + php_strval(_coalesce(run.get('usage'), 'python this_file.py "your prompt here"')))
        for x in _coalesce(run.get('extra'), []):
            L.append('  ' + x)
        hasPlaybook = False
        for n in nodes:
            if n['playbook'] is not None and not php_empty(n.get('supported')):
                hasPlaybook = True
        if hasPlaybook:
            L.append('  # Human gates (playbook nodes): on a terminal the script asks on the console;')
            L.append('  #   otherwise PLAYBOOK_GATE_MODE=auto (default: approve/acknowledge) | deny | prompt.')
        st = _coalesce(doc.get('storage'), {})
        if not php_empty(st.get('enabled')):
            folder = _coalesce(st.get('folder'), '')
            folder = folder if folder != '' else '~/Documents/synergyAI/outputs/workflow/'
            storage_line = 'ON -> ' + folder
        else:
            storage_line = 'OFF (the result is printed, not saved)'
        L.append('  # Output storage (Output node setting): ' + storage_line)
        return '\n'.join(L)

    @staticmethod
    def nodeCommentBlock(n: dict, indent: str = '') -> str:
        """PHP 728-755.

        The uniform comment block that precedes every node definition in the
        generated code (same fields in every target).
        """
        head = '# ---- node ' + str(n['id']) + ': ' + n['name'] + ' (' + n['type'] + ') '
        L = [indent + head + '-' * max(4, 78 - len(head))]

        def row(k: str, v: str) -> str:
            return indent + '#   ' + k.ljust(15) + ': ' + v

        provider = _coalesce(n.get('provider'), '')
        model = _coalesce(n.get('model'), '')
        if provider != '' or model != '':
            s = (provider if not php_empty(provider) else '?') + ' / ' + (model if model != '' else '(platform default)')
            if _coalesce(n.get('temperature')) is not None:
                s += '   temp ' + _jsonEncodeUnescapedUnicode(float(n['temperature']))
            if _coalesce(n.get('max_tokens')) is not None:
                s += '   max_tokens ' + str(php_intval(n['max_tokens']))
            s += '   thinking ' + (n['thinking'] if _coalesce(n.get('thinking')) in ('on', 'off') else 'default')
            L.append(row('provider/model', s))
        playbook = n['playbook']
        if playbook is not None:
            pb = playbook
            s = _jsonEncodeUnescapedUnicode(php_strval(_coalesce(pb.get('title'), '')))
            if _coalesce(pb.get('bound')) is not None:
                s += " -- " + php_strval(pb['bound']) + " bound / " + str(php_intval(_coalesce(pb.get('unbound'), 0))) + ' unbound action(s)'
            s += ', writes ' + ('ON' if not php_empty(pb.get('writes')) else 'OFF')
            L.append(row('playbook', s))
        else:
            L.append(row('tools', ', '.join(n['tools']) if n['tools'] else '(none)'))
            L.append(row('skills', ', '.join(n['skills']) if n['skills'] else '(none)'))
        if n['dispatch']:
            L.append(row('dispatcher', 'routes to one of: ' + ' | '.join(n['dispatch'])))
        L.append(row('parents', ' | '.join(n['parents']) if n['parents'] else '(none)'))
        L.append(row('children', ' | '.join(n['children']) if n['children'] else '(none)'))
        if php_empty(n.get('supported')):
            L.append(row('NOTE', 'not run by this target'))
        return '\n'.join(L)
