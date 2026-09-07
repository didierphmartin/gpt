"""AttachmentDispatcher parity with Services/AttachmentDispatcher.php (39-397).

The header / elision / truncation / note strings below are copied verbatim from
the PHP source (93-135, 168-208, 389-395).
"""
import base64
import pathlib

from app.services.attachment_dispatcher import AttachmentDispatcher

FIX = pathlib.Path(__file__).resolve().parent.parent / 'fixtures' / 'attachments'

HEADER = ('[The user attached the following document(s); use them as authoritative '
          'context for the message that follows.]\n\n')
SKILL_HEADER = ('[The user attached the following document(s). They are pre-staged on the '
                'script filesystem under /scratch/ — call run_skill_script with the '
                'appropriate path in argv to process them.]\n\n')
TRUNC = '\n\n[…document truncated at 100000 characters…]'


def _ensure_fixtures(d):
    """Regenerate the committed binary fixtures if any went missing."""
    d.mkdir(parents=True, exist_ok=True)
    if not (d / 'tiny.pdf').exists():
        from pypdf import PdfWriter
        from pypdf.generic import NameObject, DictionaryObject, StreamObject
        w = PdfWriter(); page = w.add_blank_page(width=200, height=200)
        content = StreamObject(); content.set_data(b"BT /F1 12 Tf 20 100 Td (Hello PDF) Tj ET")
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                                 NameObject('/Subtype'): NameObject('/Type1'),
                                 NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject(
            {NameObject('/Font'): DictionaryObject({NameObject('/F1'): w._add_object(font)})})
        page[NameObject('/Contents')] = w._add_object(content)
        with open(d / 'tiny.pdf', 'wb') as f:
            w.write(f)
    if not (d / 'tiny.docx').exists():
        import docx
        doc = docx.Document(); doc.add_heading('Title A', level=1); doc.add_paragraph('Hello DOCX')
        t = doc.add_table(rows=1, cols=2); t.cell(0, 0).text = 'c1'; t.cell(0, 1).text = 'c2'
        doc.save(d / 'tiny.docx')
    if not (d / 'tiny.xlsx').exists():
        import openpyxl
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'S1'
        ws.append(['h1', 'h2']); ws.append([1, 'x']); wb.save(d / 'tiny.xlsx')
    if not (d / 'tiny.pptx').exists():
        from pptx import Presentation
        p = Presentation(); s = p.slides.add_slide(p.slide_layouts[5])
        s.shapes.title.text = 'Slide One'; p.save(d / 'tiny.pptx')


_ensure_fixtures(FIX)


class Db:
    def __init__(self, rows): self.rows = rows; self.calls = []

    def fetch_all(self, sql, params=None):
        self.calls.append((' '.join(sql.split()), params)); return self.rows

    def fetch_one(self, *a): return None

    def fetch_column(self, *a): return []

    def execute(self, *a): return 1

    def insert(self, *a): return 1


def _row(id_, name, path, mime):
    return {'id': id_, 'user_id': 3, 'original_name': name, 'stored_path': str(path),
            'mime_type': mime,
            'size_bytes': path.stat().st_size if isinstance(path, pathlib.Path) else 0}


def test_empty_ids_short_circuit_and_sql():
    db = Db([]); d = AttachmentDispatcher(db)
    assert d.buildPrefix([], 3) == {'prefix': '', 'image_attachments': [], 'pdf_attachments': [],
                                    'notes': []} and db.calls == []
    d.buildPrefix([1, 2], 3)
    sql, params = db.calls[0]
    assert sql == ('SELECT id, original_name, stored_path, mime_type, size_bytes '
                   'FROM chat_attachments '
                   'WHERE user_id = ? AND deleted_at IS NULL AND id IN (?,?)')
    assert list(params) == [3, 1, 2]


def test_text_pdf_docx_xlsx_pptx_are_inlined_with_php_layout(tmp_path):
    txt = tmp_path / 'notes.txt'; txt.write_text('plain text body')
    rows = [_row(1, 'notes.txt', txt, 'text/plain'),
            _row(2, 'tiny.pdf', FIX / 'tiny.pdf', 'application/pdf'),
            _row(3, 'tiny.docx', FIX / 'tiny.docx',
                 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
            _row(4, 'tiny.xlsx', FIX / 'tiny.xlsx',
                 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
            _row(5, 'tiny.pptx', FIX / 'tiny.pptx',
                 'application/vnd.openxmlformats-officedocument.presentationml.presentation')]
    out = AttachmentDispatcher(Db(rows)).buildPrefix([1, 2, 3, 4, 5], 3, provider='kimi')
    p = out['prefix']
    assert p.startswith(HEADER)
    assert p.endswith('\n\n---\n\n')
    assert p.count('\n\n---\n\n') == 5          # 4 separators + the trailing one
    assert '### notes.txt\nplain text body' in p
    assert '### tiny.pdf\nHello PDF' in p
    assert '### tiny.docx\nTitle A\nHello DOCX\n| c1 | c2 |' in p
    assert '### tiny.xlsx\n## Sheet: S1\n| h1 | h2 |\n| --- | --- |\n| 1 | x |' in p
    assert '### tiny.pptx\n## Slide 1\nSlide One' in p
    assert out['image_attachments'] == [] and out['pdf_attachments'] == [] and out['notes'] == []


def test_native_pdf_and_images_are_handed_off_not_inlined(tmp_path):
    png = tmp_path / 'i.png'; png.write_bytes(b'\x89PNG\r\n\x1a\n' + b'\x00' * 16)
    rows = [_row(1, 'i.png', png, 'image/png'),
            _row(2, 'tiny.pdf', FIX / 'tiny.pdf', 'application/pdf')]
    out = AttachmentDispatcher(Db(rows)).buildPrefix([1, 2], 3, provider='claude')
    assert out['image_attachments'][0]['name'] == 'i.png'
    assert out['image_attachments'][0]['mime_type'] == 'image/png'
    assert base64.b64decode(out['image_attachments'][0]['data']).startswith(b'\x89PNG')
    assert out['pdf_attachments'][0]['name'] == 'tiny.pdf'
    assert out['pdf_attachments'][0]['mime_type'] == 'application/pdf'
    assert base64.b64decode(out['pdf_attachments'][0]['data']).startswith(b'%PDF')
    assert out['prefix'] == ''


def test_skill_mode_elides_text_with_the_php_header(tmp_path):
    txt = tmp_path / 'a.txt'; txt.write_text('x' * 50)
    out = AttachmentDispatcher(Db([_row(1, 'a.txt', txt, 'text/plain')])).buildPrefix(
        [1], 3, provider='claude', skillModeActive=True)
    assert out['prefix'] == (
        SKILL_HEADER
        + '### a.txt (50 bytes, mime: text/plain)\n'
        + '[Pre-staged at /scratch/a.txt for skill-script access. '
        + 'Pass this path to run_skill_script via argv (e.g. -i /scratch/a.txt). '
        + 'Full content NOT inlined here to save tokens.]'
        + '\n\n---\n\n'
    )
    assert 'x' * 50 not in out['prefix']


def test_truncation_marker_at_100k_chars(tmp_path):
    txt = tmp_path / 'big.txt'; txt.write_text('y' * 100_050)
    out = AttachmentDispatcher(Db([_row(1, 'big.txt', txt, 'text/plain')])).buildPrefix(
        [1], 3, provider='kimi')
    assert out['prefix'].count('y') == 100_000
    assert out['prefix'] == HEADER + '### big.txt\n' + 'y' * 100_000 + TRUNC + '\n\n---\n\n'


def test_extract_failure_becomes_note_not_exception(tmp_path):
    bad = tmp_path / 'bad.pdf'; bad.write_bytes(b'not a pdf')
    out = AttachmentDispatcher(Db([_row(1, 'bad.pdf', bad, 'application/pdf')])).buildPrefix(
        [1], 3, provider='kimi')
    assert out['notes'] == ["Could not read 'bad.pdf'; skipped."]
    assert out['prefix'] == ''


def test_missing_file_and_empty_text_notes(tmp_path):
    gone = str(tmp_path / 'gone.txt')
    empty = tmp_path / 'empty.txt'; empty.write_text('')
    rows = [_row(1, 'gone.txt', gone, 'text/plain'), _row(2, 'empty.txt', empty, 'text/plain')]
    out = AttachmentDispatcher(Db(rows)).buildPrefix([1, 2], 3, provider='kimi')
    assert out['notes'] == ["Attachment 'gone.txt' missing on disk; skipped.",
                            "Attachment 'empty.txt' contained no readable text; skipped."]
    assert out['prefix'] == ''


def test_pdf_is_extracted_when_provider_has_no_native_pdf_and_none_provider(tmp_path):
    rows = [_row(1, 'tiny.pdf', FIX / 'tiny.pdf', 'application/pdf')]
    out = AttachmentDispatcher(Db(rows)).buildPrefix([1], 3)
    assert out['pdf_attachments'] == []
    assert out['prefix'] == HEADER + '### tiny.pdf\nHello PDF' + '\n\n---\n\n'
    # provider casing is lowered before the NATIVE_PDF_PROVIDERS check (PHP 108)
    out = AttachmentDispatcher(Db(rows)).buildPrefix([1], 3, provider='Claude')
    assert out['pdf_attachments'] and out['prefix'] == ''
