"""Standalone Google ADK workflow: diamond
Auto-generated -- backend-independent. Self-contained: MCP + skills run
in this program's own Python environment.

requirements:
    pip install google-adk litellm httpx
"""
import asyncio, json, os, subprocess, sys, threading, time, traceback, urllib.request
import httpx
from typing import Any

from google.adk.agents import LlmAgent, SequentialAgent, ParallelAgent
from google.adk.models.lite_llm import LiteLlm
from google.adk.tools import FunctionTool
from google.adk.runners import Runner
from google.adk.sessions import InMemorySessionService
from google.genai import types
def _make_model(provider: str, model: str):
    """Resolve a (provider, model) pair to an ADK model.

    Gemini -> native model string; everything else -> LiteLlm. Grok/DeepSeek/Kimi
    are OpenAI-compatible endpoints (mirrors the PHP providers / LangGraph runner),
    so they route through litellm's openai/ handler with a custom api_base. API keys
    come from the environment (.env). The (provider, model) pair is resolved by the
    generator — a blank model already fell back to the provider default upstream.
    """
    p = (provider or "claude").lower()
    if not model:
        raise RuntimeError(
            f"No model for provider {p!r}. Set a model on the agent in the editor, "
            "or a default in system_llm_settings."
        )
    if p in ("gemini", "google", "google-genai"):
        return model
    if p in ("claude", "anthropic"):
        return LiteLlm(model=model if "/" in model else "anthropic/" + model)
    if p == "openai":
        return LiteLlm(model=model if "/" in model else "openai/" + model)
    if p in ("grok", "xai"):
        return LiteLlm(
            model="openai/" + model,
            api_base="https://api.x.ai/v1",
            api_key=os.environ.get("XAI_API_KEY") or os.environ.get("GROK_API_KEY"),
        )
    if p == "deepseek":
        return LiteLlm(
            model="openai/" + model,
            api_base="https://api.deepseek.com",
            api_key=os.environ.get("DEEPSEEK_API_KEY"),
        )
    if p == "kimi":
        kwargs = dict(
            model="openai/" + model,
            api_base="https://api.moonshot.ai/v1",
            api_key=os.environ.get("KIMI_API_KEY"),
        )
        if model.startswith("kimi-k2"):
            # K2 enforces non-thinking sampling; mirror the PHP KimiProvider.
            kwargs["extra_body"] = {"thinking": {"type": "disabled"}}
        return LiteLlm(**kwargs)
    # Fallback: best-effort litellm prefixed spec.
    return LiteLlm(model=model if "/" in model else p + "/" + model)
MCP_SERVERS = {}
TOOL_CATALOG = {}
def _normalize_mcp_url(url: str) -> str:
    """Ensure the URL points to the MCP endpoint.

    - If it already ends with /mcp, leave it alone.
    - If it ends with a .php file, it IS the endpoint already (XAMPP-style
      MCP server scripts like mcp-server.php). Don't append anything.
    - Otherwise, append /mcp per the MCP convention.
    """
    url = url.rstrip("/")
    if url.endswith("/mcp") or url.endswith(".php"):
        return url
    return url + "/mcp"

def _init_mcp_session(url: str, headers: dict) -> bool:
    """Initialize an MCP session with the server.

    The MCP protocol requires a handshake before tool calls:
    1. Client sends 'initialize' with protocol version and capabilities
    2. Server responds with its capabilities
    3. Client sends 'notifications/initialized' to confirm

    Returns True on success, False on failure.
    """
    mcp_url = _normalize_mcp_url(url)
    init_req = {
        "jsonrpc": "2.0", "id": 1, "method": "initialize",
        "params": {
            "protocolVersion": "2024-11-05",
            "clientInfo": {"name": "LangGraph-Workflow", "version": "1.0.0"},
            "capabilities": {}
        }
    }
    try:
        with httpx.Client(timeout=30) as c:
            r = c.post(mcp_url, json=init_req, headers=headers)
            r.raise_for_status()
            data = _parse_mcp_response(r.text)
            if data is None or "error" in (data or {}):
                return False
            # Send initialized notification
            c.post(mcp_url, json={
                "jsonrpc": "2.0",
                "method": "notifications/initialized",
                "params": {}
            }, headers=headers)
            return True
    except Exception as e:
        print(f"[warn] MCP init failed for {url}: {e}")
        return False

def _parse_mcp_response(text: str) -> dict | None:
    """Parse a JSON-RPC response from an MCP server.

    MCP servers may respond in two formats:
    - Plain JSON: standard JSON-RPC response body
    - SSE (Server-Sent Events): lines prefixed with 'data:' containing JSON
    This function tries plain JSON first, then falls back to SSE parsing.
    """
    import json as _json
    try:
        return _json.loads(text)
    except _json.JSONDecodeError:
        pass
    for line in text.split("\n"):
        line = line.strip()
        if line.startswith("data:"):
            data = line[5:].strip()
            if data:
                try:
                    return _json.loads(data)
                except _json.JSONDecodeError:
                    continue
    return None

def _call_mcp_tool(server_url: str, tool_name: str,
                   arguments: dict) -> str:
    """Call a tool on an MCP server via JSON-RPC 2.0.

    Sequence: init session -> send tools/call -> parse response -> extract text.
    The MCP response contains a 'content' array; we extract all text items
    and join them. If no text is found, falls back to raw JSON.
    """
    mcp_url = _normalize_mcp_url(server_url)
    headers = {"Content-Type": "application/json",
               "Accept": "application/json, text/event-stream, */*"}

    _init_mcp_session(server_url, headers)

    request = {
        "jsonrpc": "2.0", "id": int(time.time()),
        "method": "tools/call",
        "params": {"name": tool_name,
                   "arguments": arguments or {}}
    }
    try:
        with httpx.Client(timeout=180) as c:
            r = c.post(mcp_url, json=request, headers=headers)
            r.raise_for_status()
            data = _parse_mcp_response(r.text)
    except Exception as e:
        return json.dumps({"error": str(e)})

    if data is None:
        return json.dumps({"error": "invalid response"})
    if "error" in data:
        return json.dumps({"error": data["error"].get("message", "unknown")})

    content = (data.get("result") or {}).get("content", [])
    texts = [c.get("text", "") for c in content
             if isinstance(c, dict) and c.get("type") == "text"]
    return "\n".join(t for t in texts if t) or json.dumps(data.get("result"))[:8000]

def _convert_doc_to_markdown(path: str) -> str:
    """Read a file and return its contents as Markdown.

    Text-native formats (HTML/MD/TXT/JSON/YAML/CSV/XML/source) are returned
    verbatim. Binary office formats and PDFs are routed to their matching
    pure-Python parser. Imports are lazy so the script doesn't pull in a
    library unless the corresponding format is actually attached.

    Raises FileNotFoundError if the path doesn't exist, ValueError for an
    unsupported extension, ImportError if the format's library is missing.
    """
    from pathlib import Path
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"document not found: {path}")
    ext = p.suffix.lower()

    text_native = {
        ".html", ".htm", ".md", ".markdown", ".txt", ".text",
        ".json", ".yaml", ".yml", ".xml", ".csv", ".tsv",
        ".css", ".scss", ".less",
        ".js", ".mjs", ".ts", ".tsx", ".jsx",
        ".py", ".rb", ".go", ".rs", ".java", ".c", ".cc", ".cpp", ".h", ".hpp",
        ".sh", ".bash", ".zsh", ".sql", ".log", ".rtf", ".ini", ".toml",
    }
    if ext in text_native:
        return p.read_text(encoding="utf-8", errors="replace")

    if ext == ".docx":
        import mammoth
        with p.open("rb") as f:
            return mammoth.convert_to_markdown(f).value

    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(str(p))
        sections = [f"# {p.name}", f"_{len(prs.slides)} slide(s)_"]
        for i, slide in enumerate(prs.slides, start=1):
            layout = slide.slide_layout.name if slide.slide_layout else "?"
            sections.append(f"## Slide {i} -- {layout}")
            title_text = ""
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                title_text = slide.shapes.title.text_frame.text.strip()
            if title_text:
                sections.append(f"### {title_text}")
            body = []
            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue
                if not getattr(shape, "has_text_frame", False):
                    continue
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text or "" for run in para.runs).strip()
                    if not line and para.text:
                        line = para.text.strip()
                    if line:
                        indent = "  " * (para.level or 0)
                        body.append(f"{indent}- {line}")
            if body:
                sections.append("\n".join(body))
            if slide.has_notes_slide:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()
                if notes:
                    sections.append("### Notes")
                    sections.append(notes)
        return "\n\n".join(sections) + "\n"

    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(str(p), data_only=True)
        sections = [f"# {p.name}", f"_{len(wb.sheetnames)} sheet(s)_"]
        max_rows = 200
        for name in wb.sheetnames:
            ws = wb[name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                row_vals = list(row)
                while row_vals and (row_vals[-1] is None or row_vals[-1] == ""):
                    row_vals.pop()
                if row_vals or not rows:
                    rows.append(row_vals)
            while rows and not any(c not in (None, "") for c in rows[-1]):
                rows.pop()
            sections.append(f"## Sheet: {name}")
            sections.append(f"_{len(rows)} row(s) -- {ws.max_column} col(s)_")
            if not rows:
                sections.append("_(empty)_")
                continue
            width = max(len(r) for r in rows)
            rows = [list(r) + [""] * (width - len(r)) for r in rows]
            truncated = len(rows) > max_rows
            if truncated:
                rows = rows[:max_rows]
            def cell(v):
                return ("" if v is None else str(v)).replace("|", "\\|").replace("\n", " ")
            table = ["| " + " | ".join(cell(v) for v in rows[0]) + " |",
                     "| " + " | ".join(["---"] * width) + " |"]
            for r in rows[1:]:
                table.append("| " + " | ".join(cell(v) for v in r) + " |")
            if truncated:
                table.append(f"\n_(showing first {max_rows} rows)_")
            sections.append("\n".join(table))
        return "\n\n".join(sections) + "\n"

    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(p))
        parts = [f"# {p.name} -- {len(reader.pages)} page(s)"]
        for i, page in enumerate(reader.pages, start=1):
            try:
                t = (page.extract_text() or "").strip()
            except Exception as e:
                t = ""
            parts.append(f"## Page {i}")
            parts.append(t if t else "_(no extractable text)_")
        return "\n\n".join(parts) + "\n"

    raise ValueError(f"unsupported document format: {ext}")

def build_tools_from_catalog() -> dict:
    return {}
catalog = build_tools_from_catalog()
START_DOCUMENTS = []
node_2 = LlmAgent(
    name="node_2",
    model=_make_model("claude", "m"),
    instruction="A",
    tools=[],
    output_key="node_2",
)

node_3 = LlmAgent(
    name="node_3",
    model=_make_model("gemini", "g"),
    instruction="B",
    tools=[],
    output_key="node_3",
)
node_4 = LlmAgent(
    name="node_4",
    model=_make_model("claude", "m"),
    instruction="Consolidate the following results into the final answer.\n\n{node_2}\n{node_3}",
    tools=[],
    output_key="node_4",
)
root_agent = SequentialAgent(
    name="workflow",
    sub_agents=[
    ParallelAgent(name="layer_1", sub_agents=[node_2, node_3]),
    node_4
    ],
)
WORKFLOW_NAME = "diamond"

async def main(user_prompt: str = "GO"):
    if START_DOCUMENTS:
        doc_parts = []
        for doc in START_DOCUMENTS:
            name = doc.get("name", "Document")
            path = doc.get("path", "")
            if not path:
                doc_parts.append(f"### {name}\n\n_(no path on attachment record)_")
                continue
            try:
                md = _convert_doc_to_markdown(path)
                doc_parts.append(f"### {name}\n\n{md}")
            except Exception as e:
                doc_parts.append(f"### {name}\n\n_(conversion failed: {e})_")
        if doc_parts:
            user_prompt = (
                "## Attached Documents\n\n"
                + "\n\n---\n\n".join(doc_parts)
                + "\n\n---\n\n"
                + user_prompt
            )
    print(f"[workflow] {WORKFLOW_NAME} starting", flush=True)
    print(f"[workflow] prompt: {user_prompt[:200]!r}", flush=True)
    session_service = InMemorySessionService()
    runner = Runner(agent=root_agent, app_name="workflow", session_service=session_service)
    session = await session_service.create_session(app_name="workflow", user_id="local", state={})
    final = ""
    t0 = time.monotonic()
    seen = set()
    content = types.Content(role="user", parts=[types.Part(text=user_prompt)])
    try:
        async for event in runner.run_async(user_id="local", session_id=session.id, new_message=content):
            author = getattr(event, "author", "?")
            if author and author not in seen:
                seen.add(author)
                print(f"[node] > {author}", flush=True)
            parts = (event.content.parts if event.content else None) or []
            for p in parts:
                fc = getattr(p, "function_call", None)
                fr = getattr(p, "function_response", None)
                if fc is not None:
                    _a = str(getattr(fc, "args", ""))[:200]
                    print(f"[tool] -> {fc.name}({_a})", flush=True)
                elif fr is not None:
                    _r = str(getattr(fr, "response", ""))[:200]
                    print(f"[tool] <- {fr.name}: {_r}", flush=True)
                else:
                    txt = (getattr(p, "text", None) or "").strip()
                    if txt:
                        print(f"[{author}] {txt[:240]}", flush=True)
            if event.is_final_response() and event.content and event.content.parts:
                final = event.content.parts[0].text or final
        print(f"[workflow] done in {time.monotonic() - t0:.1f}s, {len(final)} chars", flush=True)
    except Exception:
        print("[workflow] ERROR:", flush=True)
        traceback.print_exc()
        raise
    os.makedirs("outputs", exist_ok=True)
    _head = final[:500].lower()
    _ext = "html" if ("<!doctype" in _head or "<html" in _head) else "md"
    _slug = "".join(c if c.isalnum() else "_" for c in WORKFLOW_NAME).strip("_")[:60] or "workflow"
    _ts = time.strftime("%Y%m%d-%H%M%S")
    _out = os.path.join("outputs", f"{_slug}_{_ts}.{_ext}")
    with open(_out, "w", encoding="utf-8") as _f:
        _f.write(final)
    print(f"[workflow] result saved to {os.path.abspath(_out)}", flush=True)
    print(final)
    return final

if __name__ == "__main__":
    asyncio.run(main(sys.argv[1] if len(sys.argv) > 1 else "GO"))
